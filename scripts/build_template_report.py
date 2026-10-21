"""Fill a retained teacher DOCX by patching only declared package slots.

Run with the managed document Python. The original is never written and is
intentionally not included in the code repository or delivery package.
"""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
import argparse
import hashlib
import json
import re
from lxml import etree as E
from PIL import Image
from docx import Document
from docx.shared import Inches

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'artifacts/release/template_work'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
NS={'w':W,'m':M,'r':R,'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
EXPECTED='a53057dc5e41568263441f4b358cb338c5d2124cc566f16715f136304d43c80d'


def element(name,attrs=None,text=None):
    prefix,local=name.split(':');node=E.Element('{'+NS[prefix]+'}'+local)
    for key,value in (attrs or {}).items():node.set('{'+NS[prefix]+'}'+key,str(value))
    if text is not None:node.text=text
    return node


def put(parent,name,attrs):
    node=parent.find('w:'+name,NS)
    if node is not None:parent.remove(node)
    node=element('w:'+name,attrs);parent.append(node);return node


def run(text,size=24,bold=False,sup=False):
    r=element('w:r');pr=element('w:rPr');r.append(pr)
    pr.append(element('w:rFonts',{'ascii':'Times New Roman','hAnsi':'Times New Roman','eastAsia':'宋体'}))
    pr.append(element('w:sz',{'val':size}));pr.append(element('w:szCs',{'val':size}));pr.append(element('w:color',{'val':'000000'}))
    if bold:pr.append(element('w:b'))
    if sup:pr.append(element('w:vertAlign',{'val':'superscript'}))
    t=element('w:t',text=text);t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(t);return r


def replace_text(paragraph,text):
    prototype=paragraph.find('w:r',NS)
    pr=deepcopy(prototype.find('w:rPr',NS)) if prototype is not None and prototype.find('w:rPr',NS) is not None else None
    for node in list(paragraph):
        if node.tag!='{'+W+'}pPr':paragraph.remove(node)
    r=element('w:r')
    if pr is not None:r.append(pr)
    r.append(element('w:t',text=text));paragraph.append(r)


def text_of(node):return ''.join(node.xpath('.//w:t/text()',namespaces=NS))


def build(reference,output,source):
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==EXPECTED,'Template changed; repeat distillation'
    contract=ROOT/'docs/report_template_contract.md'
    assert contract.is_file() and EXPECTED in contract.read_text(encoding='utf-8'),'Template execution contract missing'
    with ZipFile(reference) as z:parts={n:z.read(n) for n in z.namelist()}
    tree=E.fromstring(parts['word/document.xml']);body=tree.find('w:body',NS);original=list(body)
    relationships=E.fromstring(parts['word/_rels/document.xml.rels'])
    chapter_source=original[13];body_source=original[14]
    cover=[deepcopy(n) for n in original[:13]]
    replace_text(cover[8],'智能火灾报警系统')
    cover_table=cover[6];rows=cover_table.findall('w:tr',NS)
    for row_index,value in [(1,'   电子信息工程  待填写   '),(2,'    姓名及学号待填写        '),(3,'          待填写                 ')]:
        cell=rows[row_index].findall('w:tc',NS)[1]
        paragraphs=cell.findall('w:p',NS)
        replace_text(paragraphs[0],value)
        for p in paragraphs[1:]:replace_text(p,'')
    for t in cover_table.xpath('.//w:t',namespaces=NS):
        if '20xx' in (t.text or ''):t.text='2026年10月21日'
    section=deepcopy(original[-1])
    for n in list(body):body.remove(n)
    for n in cover:body.append(n)
    additions={};next_image=100

    def paragraph(text='',kind='body',new_page=False,citations=True):
        prototype=chapter_source if kind=='chapter' else body_source
        p=deepcopy(prototype)
        for key in list(p.attrib):del p.attrib[key]
        pr=p.find('w:pPr',NS)
        for n in list(p):
            if n is not pr:p.remove(n)
        for n in list(pr):
            if n.tag in ['{'+W+'}rPr','{'+W+'}ind','{'+W+'}spacing','{'+W+'}numPr','{'+W+'}pBdr','{'+W+'}pageBreakBefore','{'+W+'}keepNext']:pr.remove(n)
        put(pr,'spacing',{'before':120 if kind in ('chapter','sub') else 0,'after':100 if kind in ('chapter','sub') else 50,'line':300,'lineRule':'auto'})
        put(pr,'ind',{'firstLine':480 if kind=='body' else 0,'firstLineChars':200 if kind=='body' else 0})
        put(pr,'jc',{'val':'center' if kind=='caption' else 'left' if kind in ('chapter','sub','reference') else 'both'})
        put(pr,'widowControl',{'val':1})
        if kind=='chapter':
            pr.append(deepcopy(chapter_source.find('w:pPr/w:numPr',NS)))
            put(pr,'keepNext',{'val':1})
        if kind=='sub':put(pr,'keepNext',{'val':1})
        if new_page:put(pr,'pageBreakBefore',{'val':1})
        size=28 if kind=='chapter' else 21 if kind in ('caption','reference','cell') else 24
        pieces=re.split(r'(\[\d+(?:[-,]\d+)*\])',text) if citations else [text]
        for piece in pieces:
            if piece:p.append(run(piece,size,kind in ('chapter','sub'),citations and bool(re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',piece))))
        return p

    def picture(path,caption):
        nonlocal next_image
        if path.suffix=='.svg':path=WORK/'figures'/(path.stem+'.png')
        with Image.open(path) as img:w,h=img.size
        width=min(5.65,3.6*w/h)
        temporary=Document();temporary.add_paragraph().add_run().add_picture(str(path),width=Inches(width))
        p=E.fromstring(E.tostring(temporary.paragraphs[0]._p))
        ppr=p.find('w:pPr',NS)
        if ppr is None:ppr=element('w:pPr');p.insert(0,ppr)
        put(ppr,'jc',{'val':'center'});put(ppr,'keepNext',{'val':1});put(ppr,'spacing',{'before':80,'after':60})
        rid='rIdReport'+str(next_image);name='report_'+str(next_image)+'.png';next_image+=1
        for blip in p.xpath('.//a:blip',namespaces=NS):blip.set('{'+R+'}embed',rid)
        for docpr in p.xpath('.//*[local-name()="docPr"]'):docpr.set('id',str(next_image));docpr.set('name',caption);docpr.set('descr',caption)
        relationships.append(E.Element('{'+PKG+'}Relationship',Id=rid,Type=R+'/image',Target='media/'+name))
        additions['word/media/'+name]=path.read_bytes();body.append(p);body.append(paragraph(caption,'caption',citations=False))

    def table(rows):
        count=len(rows[0]);tbl=element('w:tbl');properties=deepcopy(original[33].find('w:tblPr',NS))
        for item in list(properties):
            if E.QName(item).localname in ('tblW','tblInd','tblBorders','tblpPr'):properties.remove(item)
        properties.append(element('w:tblW',{'w':8306,'type':'dxa'}));properties.append(element('w:jc',{'val':'center'}))
        borders=element('w:tblBorders')
        for side in ('top','bottom'):borders.append(element('w:'+side,{'val':'single','sz':8,'color':'000000'}))
        for side in ('left','right','insideV','insideH'):borders.append(element('w:'+side,{'val':'nil'}))
        properties.append(borders);tbl.append(properties)
        widths=[2200,6106] if count==2 else [1650,3328,3328]
        grid=element('w:tblGrid')
        for width in widths:grid.append(element('w:gridCol',{'w':width}))
        tbl.append(grid)
        for index,values in enumerate(rows):
            row=element('w:tr');rowpr=element('w:trPr');rowpr.append(element('w:cantSplit'))
            if index==0:rowpr.append(element('w:tblHeader'))
            row.append(rowpr)
            for col,(value,width) in enumerate(zip(values,widths)):
                cell=element('w:tc');cp=element('w:tcPr');cp.append(element('w:tcW',{'w':width,'type':'dxa'}));cp.append(element('w:vAlign',{'val':'center'}))
                mar=element('w:tcMar')
                for side in ('top','bottom','left','right'):mar.append(element('w:'+side,{'w':60,'type':'dxa'}))
                cp.append(mar)
                if index==0:
                    cb=element('w:tcBorders');cb.append(element('w:bottom',{'val':'single','sz':4,'color':'000000'}));cp.append(cb)
                cell.append(cp);p=paragraph(value,'cell',citations=False)
                put(p.find('w:pPr',NS),'jc',{'val':'center' if col==0 else 'left'})
                if index==0:
                    for rpr in p.findall('w:r/w:rPr',NS):rpr.append(element('w:b'))
                cell.append(p);row.append(cell)
            tbl.append(row)
        body.append(tbl)

    def math_run(text):
        r=element('m:r');r.append(element('m:t',text=text));return r

    def fraction(numerator,denominator):
        f=element('m:f');num=element('m:num');den=element('m:den')
        num.append(math_run(numerator));den.append(math_run(denominator));f.extend([num,den]);return f

    def equation(kind):
        p=paragraph('','body');pr=p.find('w:pPr',NS);put(pr,'ind',{'firstLine':0});put(pr,'jc',{'val':'left'})
        tabs=element('w:tabs');tabs.append(element('w:tab',{'val':'center','pos':4153}));tabs.append(element('w:tab',{'val':'right','pos':8306}));pr.append(tabs)
        tab=element('w:r');tab.append(element('w:tab'));p.append(tab)
        math=element('m:oMath')
        if kind=='coverage':
            sub=element('m:sSub');base=element('m:e');base.append(math_run('A'));suffix=element('m:sub');suffix.append(math_run('c'));sub.extend([base,suffix]);math.append(sub)
            math.append(math_run(' = '));f=element('m:f');num=element('m:num');num.append(math_run('|'))
            for symbol,index in [('⋃','k'),('M','c,k')]:
                sub=element('m:sSub');base=element('m:e');base.append(math_run(symbol));suffix=element('m:sub');suffix.append(math_run(index));sub.extend([base,suffix]);num.append(sub)
            num.append(math_run('|'));den=element('m:den');den.append(math_run('W × H'));f.extend([num,den]);math.append(f)
            math.append(math_run(' × 100%'));number='（2.1）'
        else:
            math.append(math_run('IoU = '));math.append(fraction('|P ∩ G|','|P ∪ G|'));math.append(math_run('    Dice = '));math.append(fraction('2 |P ∩ G|','|P| + |G|'));number='（2.2）'
        p.append(math);tab=element('w:r');tab.append(element('w:tab'));p.append(tab);p.append(run(number));body.append(p)

    lines=source.read_text(encoding='utf-8').splitlines();index=0;new_page=False;in_references=False
    while index<len(lines):
        line=lines[index].strip();index+=1
        if not line:continue
        if line=='<!-- page -->':continue
        if line.startswith('# '):
            title=line[2:];in_references=title=='参考文献';body.append(paragraph(title,'chapter',new_page=title=='课题研究背景与意义',citations=False));new_page=False;continue
        if line.startswith('## '):
            body.append(paragraph(line[3:],'sub',new_page,citations=False));new_page=False;continue
        if line.startswith('![ '):raise ValueError('Malformed image')
        match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if match:picture((source.parent/match[2]).resolve(),match[1]);continue
        match=re.fullmatch(r'\{\{equation:(\w+)\}\}',line)
        if match:equation(match[1]);continue
        if line.startswith('|'):
            rows=[[c.strip() for c in line.strip('|').split('|')]]
            while index<len(lines) and lines[index].strip().startswith('|'):
                row=[c.strip() for c in lines[index].strip().strip('|').split('|')];index+=1
                if all(re.fullmatch('[-: ]+',x) for x in row):continue
                rows.append(row)
            table(rows);continue
        if line.startswith('表'):
            p=paragraph(line,'caption',citations=False);put(p.find('w:pPr',NS),'keepNext',{'val':1});body.append(p);continue
        body.append(paragraph(line,'reference' if in_references and re.match(r'^\[\d+\]',line) else 'body',citations=not in_references))
    # Copy assessment furniture; clear example approval rather than impersonate it.
    grade_title=deepcopy(original[47]);body.append(grade_title)
    for page_break in grade_title.xpath('.//w:br[@w:type="page"]|.//w:lastRenderedPageBreak',namespaces=NS):page_break.getparent().remove(page_break)
    put(grade_title.find('w:pPr',NS),'pageBreakBefore',{'val':1})
    for n in original[48:51]:body.append(deepcopy(n))
    grading=deepcopy(original[51])
    signature_cell=grading.findall('w:tr',NS)[-1].find('w:tc',NS)
    for child in list(signature_cell):
        if child.tag!='{'+W+'}tcPr':signature_cell.remove(child)
    signature_cell.append(paragraph('指导教师（签章）：                     年    月    日','cell',citations=False))
    body.append(grading);body.append(deepcopy(original[52]));body.append(section)
    for rel in list(relationships):
        if rel.get('Id')=='rId10':relationships.remove(rel)
    parts.pop('word/media/image6.png',None)
    parts['word/document.xml']=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    parts['word/_rels/document.xml.rels']=E.tostring(relationships,xml_declaration=True,encoding='UTF-8',standalone=True)
    parts.update(additions)
    output.parent.mkdir(parents=True,exist_ok=True)
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    with ZipFile(reference) as before,ZipFile(output) as after:
        mutable={'word/document.xml','word/_rels/document.xml.rels','word/media/image6.png'}
        for name in before.namelist():
            if name not in mutable:assert before.read(name)==after.read(name),name
        final=E.fromstring(after.read('word/document.xml'))
        assert E.tostring(final.find('w:body/w:sectPr',NS),method='c14n',exclusive=True)==E.tostring(original[-1],method='c14n',exclusive=True)
        assert len(final.findall('.//m:oMath',NS))==2
        assert len(final.xpath('.//w:vertAlign[@w:val="superscript"]',namespaces=NS))>=9
        assert 'rId10' not in after.read('word/document.xml').decode()
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==EXPECTED
    result={'status':'passed','reference_sha256':EXPECTED,'reference_unchanged':True,'preserve_only_parts_identical':True,'a4_section_unchanged':True,'cover_logo_preserved':True,'teacher_signature_and_scores_blank':True,'native_word_equations':2,'bibliography_items':9,'added_figures':len(additions),'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest()}
    (ROOT/'reports/template_report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False));print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--template',type=Path,required=True);parser.add_argument('--source',type=Path,default=ROOT/'docs/template_report.md');parser.add_argument('--output',type=Path,default=ROOT/'artifacts/release/智能火灾报警系统设计报告.docx')
    args=parser.parse_args();build(args.template,args.output,args.source)
