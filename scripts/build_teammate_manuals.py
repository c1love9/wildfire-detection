"""Create editable teammate manuals from maintained Markdown and teaching assets.

Run using the managed document runtime, rather than the application environment.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RELEASE = ROOT / 'artifacts/release'


def font(style, size):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def rich(paragraph, text):
    for token in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not token:
            continue
        bold = token.startswith('**') and token.endswith('**')
        code = token.startswith('`') and token.endswith('`')
        run = paragraph.add_run(token[2:-2] if bold else token[1:-1] if code else token)
        run.bold = bold
        if code:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')


def table(doc, rows):
    columns = len(rows[0])
    if columns == 2:
        widths = [6.0, 10.8] if rows[0][0] in ('位置', '项目内位置') else [3.9, 12.9]
    elif columns == 3:
        widths = [3.5, 6.0, 7.3]
    elif columns == 4:
        widths = [2.9, 4.1, 3.3, 6.5]
    else:
        raise ValueError('Unsupported manual table')
    result = doc.add_table(rows=0, cols=columns)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for name, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
            border.set(qn('w:' + name), value)
        borders.append(border)
    result._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in (('top', '100'), ('bottom', '100'), ('left', '110'), ('right', '110')):
        margin = OxmlElement('w:' + edge)
        margin.set(qn('w:w'), value)
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    result._tbl.tblPr.append(margins)
    for index, values in enumerate(rows):
        if len(values) != columns:
            raise ValueError('Mismatched table cells')
        row = result.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for column, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), 'E5EDF5' if index == 0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = index == 0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or column == 0 else WD_ALIGN_PARAGRAPH.LEFT
            rich(p, value)
            for run in p.runs:
                run.font.size = Pt(9.8)
                if index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = .3
    return result


def build(source, target):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.1)
    section.header_distance = Cm(.75)
    section.footer_distance = Cm(.75)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for name, size in (('Normal', 11), ('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5),
                       ('Caption', 9.5), ('Header', 9), ('Footer', 9), ('List Bullet', 11)):
        font(doc.styles[name], size)
    body = doc.styles['Normal'].paragraph_format
    body.line_spacing = 1.23
    body.space_after = Pt(7)
    for name in ('Heading 1', 'Heading 2'):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(13 if name == 'Heading 1' else 10)
        fmt.space_after = Pt(7)
        fmt.keep_with_next = True
        fmt.keep_together = True
    doc.styles['Title'].paragraph_format.space_after = Pt(13)
    header = section.header.paragraphs[0]
    header.add_run('智能火灾报警系统  项目协作资料')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    for field, trailing in (('PAGE', ' 页  共 '), ('NUMPAGES', ' 页')):
        element = OxmlElement('w:fldSimple')
        element.set(qn('w:instr'), field)
        footer._p.append(element)
        footer.add_run(trailing)
    lines = source.read_text(encoding='utf-8').splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.left_indent = Cm(.4)
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
            continue
        if line.startswith('|'):
            rows = [[cell.strip() for cell in line.strip('|').split('|')]]
            while index < len(lines) and lines[index].strip().startswith('|'):
                candidate = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                index += 1
                if not all(re.fullmatch(r'[-: ]+', cell) for cell in candidate):
                    rows.append(candidate)
            table(doc, rows)
            continue
        match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if match:
            image = source.parent / match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
            shape = p.add_run().add_picture(str(image), width=Cm(16.4))
            shape._inline.docPr.set('descr', match.group(1))
            caption = doc.add_paragraph(match.group(1), 'Caption')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next = False
            continue
        level = next((prefix for prefix in ('### ', '## ', '# ') if line.startswith(prefix)), None)
        if level:
            p = doc.add_paragraph(line[len(level):], {'# ': 'Title', '## ': 'Heading 1', '### ': 'Heading 2'}[level])
            continue
        bullet = line.startswith('- ')
        p = doc.add_paragraph(style='List Bullet' if bullet else None)
        rich(p, line[2:] if bullet else line)
        if re.match(r'\d+\. ', line):
            p.paragraph_format.left_indent = Cm(.5)
            p.paragraph_format.first_line_indent = Cm(-.5)
    title = next(line[2:] for line in lines if line.startswith('# '))
    doc.core_properties.title = title
    doc.core_properties.subject = '项目原理 实现方法与真实人工标注操作'
    doc.core_properties.author = '项目组'
    doc.core_properties.keywords = 'fire smoke YOLO GroundingDINO SAM 人工标注'
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return {'document': target.relative_to(ROOT).as_posix(), 'source': source.relative_to(ROOT).as_posix(),
            'title': title, 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'figures': len(doc.inline_shapes)}


def main():
    records = [build(ROOT / 'docs/teammate_project_guide.md', RELEASE / '智能火灾报警系统项目入门手册.docx'),
               build(ROOT / 'docs/teammate_annotation_manual.md', RELEASE / '火焰烟雾数据标注操作手册.docx')]
    manifest = ROOT / 'artifacts/teammate_manual/build.json'
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(records, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
