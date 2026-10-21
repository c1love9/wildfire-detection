"""Build an editable engineering report from the tracked Markdown source.

Run using the managed document runtime (python-docx and Pillow required).
Generated files stay in the ignored artifacts/release directory.
"""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]


def font(style,size):
    style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0);style.font.underline=False
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')


def table(document,rows):
    count=len(rows[0]);result=document.add_table(rows=1,cols=count)
    result.alignment=WD_TABLE_ALIGNMENT.CENTER;result.autofit=False
    widths=([1.15,2.85,2.9] if count==3 else [1.55,5.35])
    for column,width in zip(result.columns,widths):column.width=Inches(width)
    properties=result._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    properties.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ('top','bottom','left','right'):
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    properties.append(margins)
    for index,values in enumerate(rows):
        cells=result.rows[0].cells if index==0 else result.add_row().cells
        row_properties=cells[0]._tc.getparent().get_or_add_trPr()
        avoid=OxmlElement('w:cantSplit');row_properties.append(avoid)
        if index==0:row_properties.append(OxmlElement('w:tblHeader'))
        for col,(cell,value,width) in enumerate(zip(cells,values,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DAE7F1' if index==0 else ('F6F8FA' if index%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(fill)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
            if col==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(value);run.bold=index==0;run.font.size=Pt(10.5)
    document.add_paragraph().paragraph_format.space_after=Pt(1)


def build(source,output):
    document=Document();section=document.sections[0]
    # The bundled starter has a blue Title paragraph border; remove all
    # decorative paragraph borders while preserving the explicit table grid.
    for border in list(document.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.7);section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.8);section.right_margin=Inches(.8)
    section.header_distance=Inches(.25);section.footer_distance=Inches(.25)
    font(document.styles['Normal'],11.5)
    document.styles['Normal'].paragraph_format.line_spacing=1.18
    document.styles['Normal'].paragraph_format.space_after=Pt(8)
    for name,size in [('Title',25),('Subtitle',13),('Heading 1',17),('Heading 2',13)]:
        font(document.styles[name],size)
        document.styles[name].paragraph_format.space_before=Pt(10)
        document.styles[name].paragraph_format.space_after=Pt(9)
    for name in ('Title','Heading 1','Heading 2'):
        document.styles[name].paragraph_format.keep_with_next=True
    font(document.styles['List Bullet'],11.5)
    font(document.styles['Header'],9);font(document.styles['Footer'],9)
    header=section.header.paragraphs[0];header.text='第四题实习项目  智能火灾报警系统';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('通用可编辑稿  ·  正式数据与实验待补     第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    lines=source.read_text(encoding='utf-8').splitlines();index=0
    while index<len(lines):
        line=lines[index].strip();index+=1
        if not line:continue
        if line=='<!-- page -->':continue
        if line.startswith('|'):
            rows=[[c.strip() for c in line.strip('|').split('|')]]
            while index<len(lines) and lines[index].strip().startswith('|'):
                row=[c.strip() for c in lines[index].strip().strip('|').split('|')];index+=1
                if all(re.fullmatch(r'[-: ]+',cell) for cell in row):continue
                rows.append(row)
            table(document,rows);continue
        image_match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if image_match:
            path=(source.parent/image_match[2]).resolve()
            if not path.is_file():raise FileNotFoundError(path)
            with Image.open(path) as img:width,height=img.size
            picture_width=min(6.6,3.65*width/height)
            p=document.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Inches(picture_width));continue
        if line.startswith('# '):document.add_paragraph(line[2:],'Title');continue
        if line.startswith('## '):document.add_paragraph(line[3:],'Heading 1');continue
        if line.startswith('### '):document.add_paragraph(line[4:],'Heading 2');continue
        p=document.add_paragraph(line[2:] if line.startswith('- ') else line,'List Bullet' if line.startswith('- ') else None)
        if line.startswith('图'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:run.font.size=Pt(9.5)
    document.core_properties.title='智能火灾报警系统设计报告'
    document.core_properties.subject='系统功能 实验工具与待补正式实验'
    document.core_properties.author='c1love9'
    output.parent.mkdir(parents=True,exist_ok=True);document.save(output)
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,default=ROOT/'docs/design_report.md');parser.add_argument('--output',type=Path,default=ROOT/'artifacts/release/智能火灾报警系统通用稿.docx')
    args=parser.parse_args();build(args.source,args.output)
